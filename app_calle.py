from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import re
import unicodedata

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from cryptography.fernet import Fernet, InvalidToken
except ImportError:  # pragma: no cover - la dependencia se valida al instalar
    Fernet = None
    InvalidToken = ValueError

from src import persat_api, siscor_db


ITINERARY_PATH = Path("data/itinerario_vendedores_calle.xlsx")
ENCRYPTED_ITINERARY_PATH = Path("data/itinerario_vendedores_calle.csv.enc")
WEEKDAY_TO_OFFSET = {
    "LUNES": 0,
    "MARTES": 1,
    "MIERCOLES": 2,
    "MIERCOLES": 2,
    "JUEVES": 3,
    "VIERNES": 4,
    "SABADO": 5,
    "DOMINGO": 6,
}


st.set_page_config(
    page_title="Efectividad de Calle",
    page_icon="BC",
    layout="wide",
)


def _normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text.strip().upper())


def _money(value: float) -> str:
    return f"$ {value:,.0f}".replace(",", ".")


def _number(value: float) -> str:
    return f"{value:,.0f}".replace(",", ".")


def _percent(value: float) -> str:
    return f"{value:.1f}%"


def _week_bounds(reference: date) -> tuple[date, date]:
    start = reference - timedelta(days=reference.weekday())
    return start, start + timedelta(days=6)


def _read_itinerary(uploaded_file) -> pd.DataFrame:
    raw = None
    if uploaded_file is not None:
        raw = pd.read_excel(uploaded_file, sheet_name=0)
    elif ENCRYPTED_ITINERARY_PATH.exists() and Fernet is not None:
        key = siscor_db._snapshot_key()
        if key:
            try:
                payload = Fernet(key.encode("utf-8")).decrypt(ENCRYPTED_ITINERARY_PATH.read_bytes())
                raw = pd.read_csv(BytesIO(payload))
            except (InvalidToken, ValueError):
                raw = None
    if raw is None and ITINERARY_PATH.exists():
        raw = pd.read_excel(ITINERARY_PATH, sheet_name=0)
    if raw is None:
        return pd.DataFrame(columns=["id_cliente", "cliente", "vendedor", "dia_visita"])

    column_map = {
        "Nro": "id_cliente",
        "Razon Social": "cliente",
        "Zona": "vendedor",
        "Dia de Visita": "dia_visita",
    }
    df = raw.rename(columns=column_map)
    for column in column_map.values():
        if column not in df.columns:
            df[column] = ""
    df = df.loc[:, ["id_cliente", "cliente", "vendedor", "dia_visita"]].copy()
    df["id_cliente"] = df["id_cliente"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    df["cliente"] = df["cliente"].fillna("").astype(str).str.strip()
    df["vendedor"] = df["vendedor"].fillna("").astype(str).str.strip().str.upper()
    df["dia_norm"] = df["dia_visita"].map(_normalize)
    df = df[df["id_cliente"].ne("") & df["vendedor"].ne("") & df["dia_norm"].isin(WEEKDAY_TO_OFFSET)].copy()
    return df.drop_duplicates(["id_cliente", "vendedor", "dia_norm"]).reset_index(drop=True)


def _expand_week(itinerary: pd.DataFrame, week_start: date) -> pd.DataFrame:
    if itinerary.empty:
        return itinerary.assign(fecha_programada=pd.NaT)
    out = itinerary.copy()
    out["fecha_programada"] = out["dia_norm"].map(lambda day: week_start + timedelta(days=WEEKDAY_TO_OFFSET[day]))
    out["dia_visita"] = out["dia_norm"].str.title()
    return out.loc[:, ["fecha_programada", "dia_visita", "vendedor", "id_cliente", "cliente"]]


def _client_filter(client_ids: list[str], alias: str = "f") -> tuple[str, tuple[str, ...]]:
    ids = [str(item).strip() for item in client_ids if str(item).strip()]
    if not ids:
        return " AND 1 = 0", ()
    placeholders = ", ".join("?" for _ in ids)
    return f" AND CAST({alias}.id_cliente AS varchar(50)) IN ({placeholders})", tuple(ids)


@st.cache_data(ttl=300, show_spinner=False)
def _current_debt(client_ids: tuple[str, ...]) -> pd.DataFrame:
    client_sql, params = _client_filter(list(client_ids), "f")
    signed_balance = siscor_db._signed_balance("f", "saldo")
    return siscor_db.read_sql(
        f"""
        SET NOCOUNT ON;
        SELECT
            CAST(f.id_cliente AS varchar(50)) AS id_cliente,
            SUM(CASE WHEN f.saldo <> 0 THEN {signed_balance} ELSE 0 END) AS saldo_actual,
            SUM(CASE WHEN f.saldo <> 0
                AND CAST(COALESCE(f.fecha_vencimiento, f.fecha) AS date) < CAST(GETDATE() AS date)
                THEN {signed_balance} ELSE 0 END) AS saldo_vencido,
            MAX(CASE WHEN f.saldo <> 0
                THEN DATEDIFF(day, COALESCE(f.fecha_vencimiento, f.fecha), GETDATE())
                ELSE 0 END) AS dias_mora
        FROM dbo.cli_factura f
        WHERE ISNULL(f.Anulado, 0) = 0
          {siscor_db._authorized_invoice_filter("f")}
          AND f.saldo <> 0
          {siscor_db._balance_document_filter("f")}
          {client_sql}
        GROUP BY f.id_cliente;
        """,
        params,
    )


@st.cache_data(ttl=300, show_spinner=False)
def _collections(client_ids: tuple[str, ...], start: str, end: str) -> pd.DataFrame:
    client_sql, params = _client_filter(list(client_ids), "f")
    signed_total = siscor_db._signed_total("f", "total")
    return siscor_db.read_sql(
        f"""
        SET NOCOUNT ON;
        SELECT
            CAST(f.id_cliente AS varchar(50)) AS id_cliente,
            SUM({signed_total}) AS cobrado_semana,
            COUNT(*) AS comprobantes_cobrados
        FROM dbo.cli_factura f
        WHERE ISNULL(f.Anulado, 0) = 0
          {siscor_db._authorized_invoice_filter("f")}
          AND ISNULL(f.cobrado, 0) = 1
          AND f.fecha_liquidacion IS NOT NULL
          AND CAST(f.fecha_liquidacion AS date) BETWEEN ? AND ?
          {siscor_db._credit_document_filter("f")}
          {client_sql}
        GROUP BY f.id_cliente;
        """,
        (start, end, *params),
    )


@st.cache_data(ttl=300, show_spinner=False)
def _sales(client_ids: tuple[str, ...], start: str, end: str) -> pd.DataFrame:
    client_sql, params = _client_filter(list(client_ids), "f")
    signed_total = siscor_db._signed_total("f", "total")
    return siscor_db.read_sql(
        f"""
        SET NOCOUNT ON;
        SELECT
            CAST(f.id_cliente AS varchar(50)) AS id_cliente,
            SUM({signed_total}) AS venta_semana,
            COUNT(*) AS comprobantes_venta
        FROM dbo.cli_factura f
        WHERE ISNULL(f.Anulado, 0) = 0
          {siscor_db._authorized_invoice_filter("f")}
          AND CAST(f.fecha AS date) BETWEEN ? AND ?
          {siscor_db._commercial_document_filter("f")}
          {client_sql}
        GROUP BY f.id_cliente;
        """,
        (start, end, *params),
    )


def _visit_summary(visits: pd.DataFrame) -> pd.DataFrame:
    if visits.empty:
        return pd.DataFrame(
            columns=["id_cliente", "fecha_visita", "visitas", "duracion_min", "hora_primera_visita", "vendedor_persat"]
        )
    base = visits.copy()
    base["id_cliente"] = base["id_cliente"].astype(str).str.strip()
    base = base[base["id_cliente"].ne("")]
    if base.empty:
        return pd.DataFrame(
            columns=["id_cliente", "fecha_visita", "visitas", "duracion_min", "hora_primera_visita", "vendedor_persat"]
        )
    device_zone = {
        device_id: zone
        for zone, device_ids in persat_api.ZONE_DEVICE_MAP.items()
        for device_id in device_ids
    }
    base["vendedor_zona"] = base["device_id"].map(device_zone).fillna(base["vendedor"])
    base["fecha_visita"] = pd.to_datetime(base["fecha"]).dt.date
    return (
        base.groupby(["id_cliente", "fecha_visita"], as_index=False)
        .agg(
            visitas=("id_cliente", "size"),
            duracion_min=("duracion_min", "sum"),
            hora_primera_visita=("fecha_hora", "min"),
            vendedor_persat=("vendedor", "first"),
            vendedor_zona=("vendedor_zona", "first"),
        )
    )


def _build_report(week_plan: pd.DataFrame, visits: pd.DataFrame, week_start: date, week_end: date) -> dict[str, pd.DataFrame]:
    client_ids = tuple(sorted(week_plan["id_cliente"].dropna().astype(str).unique()))
    debt = _current_debt(client_ids)
    collections = _collections(client_ids, week_start.isoformat(), week_end.isoformat())
    sales = _sales(client_ids, week_start.isoformat(), week_end.isoformat())
    visit_day = _visit_summary(visits)

    detail = week_plan.merge(
        visit_day,
        left_on=["id_cliente", "fecha_programada"],
        right_on=["id_cliente", "fecha_visita"],
        how="left",
    )

    weekly_visit = (
        visit_day.groupby("id_cliente", as_index=False)
        .agg(visitas_semana=("visitas", "sum"), minutos_semana=("duracion_min", "sum"), primera_visita_semana=("fecha_visita", "min"))
        if not visit_day.empty
        else pd.DataFrame(columns=["id_cliente", "visitas_semana", "minutos_semana", "primera_visita_semana"])
    )

    detail = detail.merge(weekly_visit, on="id_cliente", how="left")
    detail = detail.merge(debt, on="id_cliente", how="left")
    detail = detail.merge(collections, on="id_cliente", how="left")
    detail = detail.merge(sales, on="id_cliente", how="left")
    for column in (
        "visitas",
        "duracion_min",
        "visitas_semana",
        "minutos_semana",
        "saldo_actual",
        "saldo_vencido",
        "dias_mora",
        "cobrado_semana",
        "comprobantes_cobrados",
        "venta_semana",
        "comprobantes_venta",
    ):
        if column in detail.columns:
            detail[column] = pd.to_numeric(detail[column], errors="coerce").fillna(0)

    detail["visitado_dia"] = detail["visitas"].gt(0)
    detail["visitado_semana"] = detail["visitas_semana"].gt(0)
    detail["tiene_deuda_vencida"] = detail["saldo_vencido"].gt(0)
    detail["estado"] = "No visitado"
    detail.loc[detail["visitado_semana"], "estado"] = "Visitado otro dia"
    detail.loc[detail["visitado_dia"], "estado"] = "Cumplido"
    detail["diagnostico"] = "Sin alerta"
    detail.loc[detail["tiene_deuda_vencida"] & ~detail["visitado_semana"], "diagnostico"] = "Deuda vencida sin visita"
    detail.loc[detail["visitado_semana"] & detail["cobrado_semana"].eq(0) & detail["tiene_deuda_vencida"], "diagnostico"] = (
        "Visitado con deuda, sin cobranza liquidada"
    )

    summary_rows = []
    for vendor, group in detail.groupby("vendedor", dropna=False):
        vendor_visits = visit_day[visit_day.get("vendedor_zona", pd.Series(dtype=str)).eq(vendor)].copy()
        planned_ids = set(group["id_cliente"].dropna().astype(str).str.strip())
        visited_ids = set(vendor_visits["id_cliente"].dropna().astype(str).str.strip())
        outside_itinerary = len(visited_ids - planned_ids)
        planned = len(group)
        completed = int(group["visitado_dia"].sum())
        visited_week = int(group["visitado_semana"].sum())
        overdue = group[group["tiene_deuda_vencida"]].copy()
        overdue_count = len(overdue)
        overdue_visited = int(overdue["visitado_semana"].sum()) if overdue_count else 0
        overdue_amount = float(overdue["saldo_vencido"].sum()) if overdue_count else 0.0
        overdue_visited_amount = float(overdue.loc[overdue["visitado_semana"], "saldo_vencido"].sum()) if overdue_count else 0.0
        collected = float(group["cobrado_semana"].sum())
        sold = float(group["venta_semana"].sum())
        collected_clients = int(group["cobrado_semana"].gt(0).sum())
        sold_after_visit = int((group["visitado_semana"] & group["venta_semana"].gt(0)).sum())

        route_score = completed / planned * 100 if planned else 0.0
        debt_focus_score = overdue_visited / overdue_count * 100 if overdue_count else 100.0
        debt_amount_score = overdue_visited_amount / overdue_amount * 100 if overdue_amount else 100.0
        collection_score = min(collected / overdue_amount * 100, 100.0) if overdue_amount else 100.0
        effectiveness = route_score * 0.35 + debt_focus_score * 0.25 + debt_amount_score * 0.25 + collection_score * 0.15
        if effectiveness >= 80:
            diagnosis = "Muy buena semana"
        elif effectiveness >= 60:
            diagnosis = "Correcto, falta foco"
        else:
            diagnosis = "Semana debil"

        summary_rows.append(
            {
                "vendedor": vendor,
                "clientes_planificados": planned,
                "cumplidos_dia": completed,
                "visitados_semana": visited_week,
                "cumplimiento_ruta_pct": route_score,
                "clientes_con_deuda_vencida": overdue_count,
                "clientes_vencidos_visitados": overdue_visited,
                "visitas_fuera_itinerario": outside_itinerary,
                "deuda_vencida_actual": overdue_amount,
                "deuda_vencida_visitada": overdue_visited_amount,
                "cobrado_semana": collected,
                "clientes_con_cobranza": collected_clients,
                "clientes_compraron_visitados": sold_after_visit,
                "venta_semana": sold,
                "indice_efectividad": effectiveness,
                "diagnostico": diagnosis,
            }
        )

    summary = pd.DataFrame(summary_rows).sort_values("indice_efectividad", ascending=False)
    critical = detail[detail["tiene_deuda_vencida"] & ~detail["visitado_semana"]].copy()
    critical = critical.sort_values(["saldo_vencido", "dias_mora"], ascending=[False, False])
    return {"summary": summary, "detail": detail, "critical": critical}


def _to_excel(report: dict[str, pd.DataFrame], week_start: date, week_end: date) -> bytes:
    def excel_ready(df: pd.DataFrame) -> pd.DataFrame:
        out = df.copy()
        for column in out.columns:
            if isinstance(out[column].dtype, pd.DatetimeTZDtype):
                out[column] = out[column].dt.tz_localize(None)
        return out

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        excel_ready(report["summary"]).to_excel(writer, sheet_name="Resumen vendedores", index=False)
        excel_ready(report["detail"]).to_excel(writer, sheet_name="Itinerario vs visitas", index=False)
        excel_ready(report["critical"]).to_excel(writer, sheet_name="Criticos no visitados", index=False)
        pd.DataFrame(
            [
                {"campo": "Semana desde", "valor": week_start.isoformat()},
                {"campo": "Semana hasta", "valor": week_end.isoformat()},
                {"campo": "Nota", "valor": "SisCor solo lectura. Deuda vencida corresponde al saldo actual consultado."},
            ]
        ).to_excel(writer, sheet_name="Notas", index=False)

        workbook = writer.book
        executive = workbook.create_sheet("Resumen ejecutivo", 0)
        _write_executive_summary(executive, report["summary"], week_start, week_end)
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center")
            for column in sheet.columns:
                max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
                sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 2, 12), 48)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "#,##0.00"
    return output.getvalue()


def _excel_money(value: object) -> str:
    try:
        return _money(float(value or 0))
    except (TypeError, ValueError):
        return "$ 0"


def _write_executive_summary(sheet, summary: pd.DataFrame, week_start: date, week_end: date) -> None:
    title_fill = PatternFill("solid", fgColor="1F4E78")
    block_fill = PatternFill("solid", fgColor="D9EAF7")
    line = Side(style="thin", color="A6A6A6")
    border = Border(bottom=line)

    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 4
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 70
    sheet.column_dimensions["D"].width = 18

    row = 1
    for _, item in summary.sort_values("vendedor").iterrows():
        vendor = str(item.get("vendedor", "")).title()
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        cell = sheet.cell(row=row, column=2, value=f"Vendedor: {vendor}")
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal="center")
        row += 1

        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        cell = sheet.cell(row=row, column=2, value=f"Semana: {week_start:%d/%m/%Y} a {week_end:%d/%m/%Y}")
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal="center")
        row += 2

        blocks = [
            (
                "Foco asignado",
                [
                    f"{int(item.get('clientes_planificados', 0))} clientes",
                    f"{_excel_money(item.get('deuda_vencida_actual', 0))} de deuda vencida asociada",
                    f"{int(item.get('clientes_con_deuda_vencida', 0))} clientes criticos",
                    f"{int(item.get('clientes_vencidos_visitados', 0))} clientes criticos visitados",
                    f"{int(item.get('visitas_fuera_itinerario', 0))} visitas fuera de itinerario",
                ],
            ),
            (
                "Resultados",
                [
                    f"{_excel_money(item.get('cobrado_semana', 0))} regularizados",
                    f"{int(item.get('clientes_con_cobranza', 0))} clientes con cobranza liquidada",
                    f"{int(item.get('clientes_compraron_visitados', 0))} clientes compraron luego de la visita",
                    f"{_excel_money(item.get('venta_semana', 0))} de venta semanal en clientes del itinerario",
                ],
            ),
            (
                "Ejecucion",
                [
                    f"Visito {int(item.get('visitados_semana', 0))} de {int(item.get('clientes_planificados', 0))} clientes",
                    f"{int(item.get('cumplidos_dia', 0))} visitas cumplidas el dia asignado",
                    f"{float(item.get('cumplimiento_ruta_pct', 0)):.1f}% de cumplimiento del dia programado",
                    f"Indice de efectividad: {float(item.get('indice_efectividad', 0)):.1f}",
                    f"Diagnostico: {item.get('diagnostico', '')}",
                ],
            ),
        ]

        for title, lines in blocks:
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            cell = sheet.cell(row=row, column=2, value=title)
            cell.font = Font(bold=True, color="1F4E78")
            cell.fill = block_fill
            row += 1
            for text in lines:
                sheet.cell(row=row, column=2, value="-")
                sheet.cell(row=row, column=3, value=text)
                sheet.cell(row=row, column=3).alignment = Alignment(wrap_text=True)
                row += 1
            for column in range(2, 5):
                sheet.cell(row=row, column=column).border = border
            row += 2


st.title("Efectividad semanal de calle")

with st.sidebar:
    st.header("Segmentadores")
    reference = st.date_input("Semana", value=date.today())
    week_start, week_end = _week_bounds(reference)
    st.caption(f"{week_start:%d/%m/%Y} al {week_end:%d/%m/%Y}")
    uploaded = st.file_uploader("Itinerario Excel", type=["xlsx"])

itinerary = _read_itinerary(uploaded)
week_plan = _expand_week(itinerary, week_start)

if week_plan.empty:
    st.warning("No hay itinerario cargado.")
    st.stop()

vendors = sorted(week_plan["vendedor"].dropna().unique())
selected_vendors = st.sidebar.multiselect("Vendedores", vendors, default=vendors)
week_plan = week_plan[week_plan["vendedor"].isin(selected_vendors)].copy()

persat_result = persat_api.activity(week_start.isoformat(), week_end.isoformat(), tuple(selected_vendors))
if not persat_result.enabled:
    st.error(persat_result.message)
    st.stop()

report = _build_report(week_plan, persat_result.visits, week_start, week_end)
summary = report["summary"]
detail = report["detail"]
critical = report["critical"]

c1, c2, c3, c4, c5 = st.columns(5)
c1.metric("Planificados", _number(len(detail)))
c2.metric("Cumplidos dia", _number(float(detail["visitado_dia"].sum())))
c3.metric("Visitados semana", _number(float(detail["visitado_semana"].sum())))
c4.metric("Deuda vencida actual", _money(float(detail["saldo_vencido"].sum())))
c5.metric("Cobrado semana", _money(float(detail["cobrado_semana"].sum())))

st.subheader("Resumen por vendedor")
st.dataframe(
    summary,
    use_container_width=True,
    hide_index=True,
    column_config={
        "cumplimiento_ruta_pct": st.column_config.NumberColumn("Cumplimiento ruta", format="%.1f%%"),
        "indice_efectividad": st.column_config.NumberColumn("Indice", format="%.1f"),
        "deuda_vencida_actual": st.column_config.NumberColumn("Deuda vencida actual", format="$ %.0f"),
        "deuda_vencida_visitada": st.column_config.NumberColumn("Deuda vencida visitada", format="$ %.0f"),
        "cobrado_semana": st.column_config.NumberColumn("Cobrado semana", format="$ %.0f"),
        "venta_semana": st.column_config.NumberColumn("Venta semana", format="$ %.0f"),
    },
)

chart_cols = st.columns(2)
with chart_cols[0]:
    fig = px.bar(summary.sort_values("indice_efectividad"), x="indice_efectividad", y="vendedor", orientation="h")
    fig.update_layout(xaxis_title="Indice", yaxis_title="", height=360)
    st.plotly_chart(fig, use_container_width=True)
with chart_cols[1]:
    fig = px.bar(summary.sort_values("cumplimiento_ruta_pct"), x="cumplimiento_ruta_pct", y="vendedor", orientation="h")
    fig.update_layout(xaxis_title="Cumplimiento ruta %", yaxis_title="", height=360)
    st.plotly_chart(fig, use_container_width=True)

tab_detail, tab_debt, tab_visits = st.tabs(["Itinerario vs visitas", "Criticos no visitados", "Visitas reales"])
with tab_detail:
    st.dataframe(
        detail.sort_values(["fecha_programada", "vendedor", "cliente"]),
        use_container_width=True,
        hide_index=True,
    )
with tab_debt:
    st.dataframe(
        critical.loc[
            :,
            [
                "fecha_programada",
                "vendedor",
                "id_cliente",
                "cliente",
                "saldo_vencido",
                "dias_mora",
                "diagnostico",
            ],
        ],
        use_container_width=True,
        hide_index=True,
        column_config={
            "saldo_vencido": st.column_config.NumberColumn("Deuda vencida", format="$ %.0f"),
            "dias_mora": st.column_config.NumberColumn("Dias mora", format="%.0f"),
        },
    )
with tab_visits:
    st.dataframe(persat_result.visits, use_container_width=True, hide_index=True)

excel_bytes = _to_excel(report, week_start, week_end)
st.download_button(
    "Exportar informe semanal",
    data=excel_bytes,
    file_name=f"informe_efectividad_calle_{week_start:%Y%m%d}_{week_end:%Y%m%d}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
)
