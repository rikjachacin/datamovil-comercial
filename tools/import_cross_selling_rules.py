from __future__ import annotations

from io import StringIO
from pathlib import Path
import sys

import pandas as pd
from cryptography.fernet import Fernet
from openpyxl import load_workbook


OUTPUT_PATH = Path("data/cross_selling_rules.csv.enc")
KEY_PATH = Path("data/snapshot.key")
SHEET_NAME = "Cruces sugeridos"
LOGICAL_SHEET_NAME = "Hipótesis lógicas"
ALLOWED_DECISIONS = {
    "Piloto prioritario",
    "Piloto supervisado",
    "Piloto segmentado",
}
LOGICAL_RULES = {
    (22380, 22383): "",
    (22380, 22283): "",
    (22383, 22380): "-FOODIE",
    (22381, 22327): "PALITA|PALA.*(?:SANIT|HIGIEN)",
    (22346, 22327): "PALITA|PALA.*(?:SANIT|HIGIEN)",
    (22404, 22306): "",
    (22371, 22393): "BOLSA|BOLSITA",
    (22406, 22286): "GATO|FELIN",
}


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Uso: python tools/import_cross_selling_rules.py reglas.xlsx")

    source = Path(sys.argv[1])
    if not source.exists():
        raise FileNotFoundError(f"No se encontro {source}")
    if not KEY_PATH.exists():
        raise FileNotFoundError("No se encontro data/snapshot.key")

    workbook = load_workbook(source, read_only=False, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Falta la hoja {SHEET_NAME}")

    sheet = workbook[SHEET_NAME]
    rows = list(sheet.iter_rows(min_row=6, values_only=True))
    records: list[dict[str, object]] = []
    for row in rows:
        if not row or row[0] is None or str(row[4]).strip() not in ALLOWED_DECISIONS:
            continue
        records.append(
            {
                "origen_id": int(row[0]),
                "origen": str(row[1]).strip(),
                "destino_id": int(row[2]),
                "destino": str(row[3]).strip(),
                "decision": str(row[4]).strip(),
                "segmento": str(row[5]).strip(),
                "coincidencia": float(row[8] or 0),
                "sugerencia": str(row[9]).strip(),
                "validacion": str(row[10]).strip(),
                "filtro_producto": "",
            }
        )

    if LOGICAL_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Falta la hoja {LOGICAL_SHEET_NAME}")
    logical_sheet = workbook[LOGICAL_SHEET_NAME]
    for row in logical_sheet.iter_rows(min_row=6, values_only=True):
        if not row or row[0] is None:
            continue
        pair = (int(row[0]), int(row[2]))
        if pair not in LOGICAL_RULES:
            continue
        records.append(
            {
                "origen_id": pair[0],
                "origen": str(row[1]).strip(),
                "destino_id": pair[1],
                "destino": str(row[3]).strip(),
                "decision": str(row[4]).strip(),
                "segmento": str(row[5]).strip(),
                "coincidencia": 0,
                "sugerencia": str(row[9]).strip(),
                "validacion": str(row[7]).strip(),
                "filtro_producto": LOGICAL_RULES[pair],
            }
        )

    if not records:
        raise ValueError("No se encontraron reglas habilitadas.")

    buffer = StringIO()
    pd.DataFrame(records).to_csv(buffer, index=False)
    encrypted = Fernet(KEY_PATH.read_bytes().strip()).encrypt(buffer.getvalue().encode("utf-8"))
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_bytes(encrypted)
    print(f"{OUTPUT_PATH}: {len(records)} reglas")


if __name__ == "__main__":
    main()
